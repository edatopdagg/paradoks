
from __future__ import annotations

import csv
import io
import json
import re
import os
import shutil

import pymupdf
import pytesseract
import xlrd
from PIL import Image, ImageOps
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
from pptx import Presentation

from app.services.embedding_service import EmbeddingService


SUPPORTED_EXTENSIONS = {
    ".xlsx",
    ".xls",
    ".csv",
    ".pdf",
    ".docx",
    ".pptx",
    ".txt",
    ".md",
    ".json",
    ".xml",
    ".png",
    ".jpg",
    ".jpeg",
    ".webp",
    ".bmp",
    ".tif",
    ".tiff",
}


@dataclass
class ExtractedRow:
    row_number: int
    values: dict[str, str]
    text: str

    source_filename: str = ""
    source_page: int | None = None
    source_kind: str = ""


def _clean(value: Any) -> str:
    if value is None:
        return ""

    return re.sub(
        r"\s+",
        " ",
        str(value),
    ).strip()


def _meaningful(values: list[str]) -> bool:
    joined = " ".join(values).strip()

    if not joined:
        return False

    # Salt s?ra numaras? gibi sat?rlar? at.
    if re.fullmatch(
        r"[\d.\-]+",
        joined,
    ):
        return False

    return True



def _normalized_match_text(
    value: str,
) -> str:
    text = _clean(value).casefold()

    translation = str.maketrans(
        {
            "\u0131": "i",
            "\u015f": "s",
            "\u011f": "g",
            "\u00fc": "u",
            "\u00f6": "o",
            "\u00e7": "c",
        }
    )
    text = text.translate(translation)

    # Türkçe büyük İ'nin casefold sonrası oluşturabildiği combining dot'u temizle.
    text = text.replace("\u0307", "")

    # Excel/Word kaynaklı Unicode tirelerini tek forma indir.
    for dash in (
        "\u2010",
        "\u2011",
        "\u2012",
        "\u2013",
        "\u2014",
        "\u2212",
    ):
        text = text.replace(dash, "-")

    text = re.sub(
        r"\s*-\s*",
        "-",
        text,
    )

    text = re.sub(
        r"\b24\s*[x/]\s*7\b",
        "7/24",
        text,
    )
    text = re.sub(
        r"\b7\s*x\s*24\b",
        "7/24",
        text,
    )

    return re.sub(
        r"\s+",
        " ",
        text,
    ).strip()




def _header_score(
    values: list[str],
) -> int:

    text = " | ".join(
        _normalized_match_text(
            value
        )
        for value
        in values
        if value
    )

    weights = (
        (
            "teknik gereksinim",
            8,
        ),
        (
            "requirement",
            5,
        ),
        (
            "compliance status",
            5,
        ),
        (
            "category",
            3,
        ),
        (
            "alan",
            2,
        ),
        (
            "madde",
            2,
        ),
        (
            "uygunluk",
            3,
        ),
        (
            "tedarikci yaniti",
            3,
        ),
    )

    return sum(
        weight
        for marker, weight
        in weights
        if marker in text
    )


def _unique_headers(
    values: list[str],
) -> list[str]:

    headers: list[str] = []
    seen: dict[str, int] = {}

    for index, value in enumerate(
        values
    ):

        base = (
            _clean(value)
            or f"column_{index + 1}"
        )

        count = seen.get(
            base,
            0,
        )

        seen[
            base
        ] = count + 1

        if count:
            headers.append(
                f"{base}_{count + 1}"
            )
        else:
            headers.append(
                base
            )

    return headers


def _company_support_status(
    row: ExtractedRow,
) -> str | None:

    text = _normalized_match_text(
        row.text
    )

    # -----------------------------------------------------
    # NEGATIVE / NOT SUPPORTED
    # -----------------------------------------------------
    #
    # Negatif kontrol daima once gelir.
    # "desteklenmemektedir" icinde olumlu bir kelime
    # parcasi bulunmasi yanlis FC uretmemeli.
    # -----------------------------------------------------

    negative = (
        "not supported",
        "unsupported",
        "non-compliance",
        "non compliance",
        "uygun degil",
        "desteklenmiyor",
        "desteklenmez",
        "desteklenmemektedir",
        "destegi bulunmamaktadir",
        "destek bulunmamaktadir",
        "bulunmamaktadir",
        "mevcut degildir",
        "saglanmamaktadir",
        "saglanmaz",
        "kullanilamamaktadir",
        "kullanilamaz",
    )

    if any(
        marker in text
        for marker in negative
    ):
        return "NC"

    # -----------------------------------------------------
    # PARTIAL
    # -----------------------------------------------------

    partial = (
        "partly supported",
        "partially supported",
        "partial compliance",
        "kismen uygun",
        "kismi uygun",
        "kismen desteklenmektedir",
        "kismen desteklenir",
        "kismen destekli",
        "sinirli destek",
        "limited support",
    )

    if any(
        marker in text
        for marker in partial
    ):
        return "PC"

    # -----------------------------------------------------
    # FULL / POSITIVE CAPABILITY ASSERTION
    # -----------------------------------------------------
    #
    # Sirket ozellik dokumanlarinda "fully supported"
    # yazmasi beklenemez. Dogal teknik beyanlar:
    #
    #   IPv6 destegi bulunmaktadir.
    #   HTTP/2 desteklenmektedir.
    #   AES-256 mevcuttur.
    #
    # Bunlar acik pozitif capability kanitidir.
    # -----------------------------------------------------

    full = (
        "fully supported",
        "full compliance",
        "tam uygun",
        "tam karsilaniyor",

        "desteklenmektedir",
        "desteklenir",
        "desteklidir",
        "destegi bulunmaktadir",
        "destek bulunmaktadir",
        "bulunmaktadir",
        "mevcuttur",
        "saglanmaktadir",
        "saglanir",
        "kullanilabilmektedir",
        "kullanilabilir",

        "is supported",
        "are supported",
        "supports",
        "is available",
        "are available",
        "provides",
        "capable of",
        "has support",
    )

    if any(
        marker in text
        for marker in full
    ):
        return "FC"

    return None


def _document_ids(
    value: str,
) -> set[str]:

    text = _normalized_match_text(
        value
    )

    patterns = (
        r"\b3gpp\s+(?:ts|tr)\s*\d{2}[.\s_-]*\d{3}\b",
        r"\betsi\s+(?:ts|en|tr|gsm)\s*\d{2,3}(?:[.\s_-]*\d{2,3})?\b",
        r"\brfc\s*\d+\b",
        r"\biso(?:/iec)?\s*\d+(?::\d+)?\b",
    )

    found: set[str] = set()

    for pattern in patterns:

        for match in re.findall(
            pattern,
            text,
            flags=re.IGNORECASE,
        ):

            found.add(
                re.sub(
                    r"[^a-z0-9]+",
                    "",
                    match.casefold(),
                )
            )

    return found


def _technical_feature_terms(
    value: str,
) -> set[str]:

    text = _normalized_match_text(
        value
    )

    patterns = (
        ("cbc", r"\bcbc\b"),
        ("cbe", r"\bcbe\b"),
        ("cbsp", r"\bcbsp\b"),
        ("sabp", r"\bsabp\b"),
        ("sbc", r"\bsbc\b"),
        ("s1ap", r"\bs1ap\b"),
        ("n50", r"\bn50\b"),
        ("amf", r"\bamf\b"),
        ("mme", r"\bmme\b"),
        ("bsc", r"\bbsc\b"),
        ("rnc", r"\brnc\b"),
        ("enb", r"\benb\b|\benodeb\b"),
        ("gnb", r"\bgnb\b"),
        ("cap", r"\bcap\b"),
        ("ldap", r"\bldap\b"),
        (
            "active_directory",
            r"\bactive directory\b",
        ),
                (
            "rbac",
            r"\brbac\b"
            r"|\brole[- ]based access control\b"
            r"|\brol[- ]yetki\b"
            r"|\brol ve yetki\b"
            r"|\brol yetki matrisi\b"
            r"|\byetki matrisi\b",
        ),
        (
            "account_lockout",
            r"\baccount lockout\b"
            r"|\baccount lock\b"
            r"|\bhesap kilit\w*\b"
            r"|\bkullanici hesab\w*.*kilit\w*\b"
            r"|\bbasarisiz.*giris.*kilit\w*\b"
            r"|\bhatali giris.*kilit\w*\b",
        ),
        (
            "password_policy",
            r"\bpassword policy\b"
            r"|\bparola politika\w*\b"
            r"|\bpassword complexity\b"
            r"|\bparola karmasik\w*\b"
            r"|\bparola gecerlilik\b"
            r"|\bparola gecmisi\b",
        ),
        ("syslog", r"\bsyslog\b"),
        (
            "siem_integration",
            r"\bsiem\b[^|.;]*(?:integration|integrable|compliance|uyum|entegrasyon)\w*"
            r"|\b(?:integration|integrable|compliance|uyum|entegrasyon)\w*[^|.;]*\bsiem\b"
            r"|\bsiem system\b",
        ),
        ("snmpv3", r"\bsnmp\s*v?3\b|\bsnmpv3\b"),
        (
            "rest_api",
            r"\brest\b|\brestful\b",
        ),
        ("vm", r"\bvm\b|\bvirtual machine"),
        ("x86", r"\bx86\b"),
        ("vmware", r"\bvmware\b"),
        ("nutanix", r"\bnutanix\b"),
        (
            "active_active",
            r"\bactive[- ]active\b",
        ),
        (
            "active_standby",
            r"\bactive[- ]standby\b",
        ),
        (
            "geo_redundancy",
            r"\bgeo[- ]redundan\w*\b|\bcografi yedeklilik\b",
        ),
        (
            "failover",
            r"\bfailover\b|\byuk devret\w*\b",
        ),
        (
            "replication",
            r"\breplication\b|\breplikasyon\b",
        ),
        (
            "backup",
            r"\bbackup\b|\byedekle\w*\b",
        ),
        (
            "restore",
            r"\brestore\b|\bgeri yukle\w*\b",
        ),
        (
            "tls",
            r"\btls\b",
        ),
        (
            "tls_12",
            r"\btls\s*v?1[.]2\b",
        ),
        (
            "tls_13",
            r"\btls\s*v?1[.]3\b",
        ),
        (
            "mutual_tls",
            r"\bmutual tls\b|\bmtls\b|\bkarsilikli tls\b",
        ),
        ("https", r"\bhttps\b"),
        ("sftp", r"\bsftp\b"),
        (
            "unicode",
            r"\bunicode\b|\bucs[- ]?2\b",
        ),
        (
            "gsm7",
            r"\bgsm\s*7[- ]?bit\b",
        ),
        ("ecgi", r"\becgi\b"),
        (
            "tracking_area",
            r"\btracking area\b|\bizleme alani\b",
        ),
        (
            "nr_cell_id",
            r"\bnr cell id\b|\bnr hucre kimligi\b",
        ),
        ("pws", r"\bpws\b|public warning system|kamu uyari"),
        ("dbgf", r"\bdbgf\b|device[- ]based geo"),
        (
            "ack_nack",
            r"\back/nack\b|\back\b.*\bnack\b",
        ),
        (
            "xml_schema",
            r"\bxml schema\b|\bxml sema",
        ),
        ("siem", r"\bsiem\b"),
        (
            "support_24x7",
            r"\b7/24\b|\b24 hour\b.*\b7 day\b",
        ),
        (
            "availability",
            r"\bavailability\b|\berisilebilirlik\b",
        ),
        (
            "scalability",
            r"\bscalab\w*\b|\bolceklen\w*\b",
        ),
        (
            "gui",
            r"\bgui\b|grafik kullanici arabirimi",
        ),
        (
            "audit_log",
            r"\baudit log\b|\bdenetim gunluk",
        ),
        (
            "certificate",
            r"\bcertificate\b|\bsertifika\b",
        ),
        (
            "digital_signature",
            r"\bdigital signature\b|\bdijital imza\b",
        ),
        (
            "segmentation",
            r"\bsegmentation\b|\bbolumle\w*\b",
        ),
        (
            "schedule",
            r"\bschedul\w*\b|\bzamanla\w*\b",
        ),
        (
            "repeat",
            r"\brepeat\w*\b|\btekrar\w*\b",
        ),
        (
            "map_gis",
            r"\bgis\b|\bharita\b|\bmap\b",
        ),
        (
            "priority",
            r"\bpriorit\w*\b|\boncelik\w*\b",
        ),
        (
            "multi_vendor",
            r"\bmulti[- ]vendor\b|\bcok ureticili\b",
        ),
        (
            "multi_language",
            r"\bmulti[- ]language\b|\bcok dilli\b",
        ),
        (
            "time_sync",
            r"\btime synchroni[sz]\w*\b|\bzaman senkron\w*\b|\bntp\b",
        ),
        (
            "cancel",
            r"\bcancel\w*\b|\biptal\w*\b",
        ),
        (
            "update",
            r"\bupdate\w*\b|\bguncelle\w*\b",
        ),
    )

    found: set[str] = set()

    for name, pattern in patterns:

        if re.search(
            pattern,
            text,
            flags=re.IGNORECASE,
        ):
            found.add(
                name
            )

    return found


def _keyword_tokens(
    value: str,
) -> set[str]:

    text = _normalized_match_text(
        value
    )

    stopwords = {
        "cozum",
        "sistem",
        "tedarikci",
        "destek",
        "desteklenmelidir",
        "desteklenmelidir.",
        "supported",
        "fully",
        "partly",
        "must",
        "shall",
        "should",
        "icin",
        "olan",
        "olarak",
        "ilgili",
        "tum",
        "bir",
        "ve",
        "veya",
        "ile",
        "the",
        "and",
        "for",
        "all",
        "solution",
        "vendor",
        "system",
        "requirement",
        "compliance",
        "status",
        "general",
        "genel",
        "technical",
        "teknik",
    }

    tokens = set(
        re.findall(
            r"[a-z0-9][a-z0-9.+/_-]{1,}",
            text,
        )
    )

    return {
        token
        for token in tokens
        if (
            token not in stopwords
            and (
                len(token) >= 3
                or any(
                    char.isdigit()
                    for char in token
                )
            )
        )
    }


def _critical_values(
    value: str,
) -> set[str]:
    """Karşılaştırmada gerçekten anlam taşıyan nicel koşulları çıkarır."""
    text = _normalized_match_text(value)
    found: set[str] = set()

    # Yüzdeler: hem %99,999 hem 99.999% biçimini destekle.
    percent_values = re.findall(
        r"(?:%\s*(\d+(?:[.,]\d+)?)|(\d+(?:[.,]\d+)?)\s*%)",
        text,
    )
    for left, right in percent_values:
        raw = left or right
        normalized = raw.replace(",", ".")
        try:
            normalized = format(float(normalized), ".12g")
        except ValueError:
            pass
        found.add(f"percent:{normalized}")

    # Yalnız açık birim taşıyan nicel şartları al. Standart kodlarındaki
    # 36.413, 29.168 gibi sayıları nicel gereksinim sanma.
    quantity_units = {
        "million": "million",
        "milyon": "million",
        "year": "year",
        "years": "year",
        "yil": "year",
        "day": "day",
        "days": "day",
        "gun": "day",
        "hour": "hour",
        "hours": "hour",
        "saat": "hour",
        "minute": "minute",
        "minutes": "minute",
        "dakika": "minute",
        "second": "second",
        "seconds": "second",
        "saniye": "second",
        "character": "character",
        "characters": "character",
        "karakter": "character",
        "attempt": "attempt",
        "attempts": "attempt",
        "deneme": "attempt",
        "location": "location",
        "locations": "location",
        "lokasyon": "location",
        "site": "site",
        "sites": "site",
        "saha": "site",
    }
    unit_pattern = "|".join(
        sorted(quantity_units, key=len, reverse=True)
    )
    for number, unit in re.findall(
        rf"\b(\d+(?:[.,]\d+)?)\s*({unit_pattern})\b",
        text,
    ):
        found.add(
            f"quantity:{number.replace(',', '.')}:{quantity_units[unit]}"
        )

    # Açık sürüm ifadesi (v1.2 / version 1.2). Teknik standart kodları hariç.
    for version in re.findall(
        r"\b(?:v|version|surum)\s*([0-9]+(?:[.]\d+)+)\b",
        text,
    ):
        found.add(f"version:{version}")

    return found





def _critical_values_satisfied(
    required_values: set[str],
    actual_values: set[str],
) -> bool:
    """
    Kritik nicel ko?ullar?n kar??lan?p kar??lanmad???n? kontrol eder.

    ?? kural?:
    - Y?zde bazl? availability/SLA/performance de?erlerinde
      2.0 y?zde puan?na kadar eksik kabul edilir.
    - ?irket de?eri istenenden y?ksekse zaten kar??lan?r.
    - S?r?m, adet, s?re, karakter say?s? vb. di?er kritik
      de?erlerde birebir e?le?me korunur.
    """

    if not required_values:
        return True

    required_percentages: list[float] = []
    actual_percentages: list[float] = []

    required_exact: set[str] = set()
    actual_exact: set[str] = set()

    for value in required_values:

        if value.startswith("percent:"):

            try:
                required_percentages.append(
                    float(
                        value.split(
                            ":",
                            1,
                        )[1]
                    )
                )
            except ValueError:
                required_exact.add(
                    value
                )

        else:
            required_exact.add(
                value
            )

    for value in actual_values:

        if value.startswith("percent:"):

            try:
                actual_percentages.append(
                    float(
                        value.split(
                            ":",
                            1,
                        )[1]
                    )
                )
            except ValueError:
                actual_exact.add(
                    value
                )

        else:
            actual_exact.add(
                value
            )

    # Version, adet, s?re, parola uzunlu?u vb.
    # kritik de?erler toleranss?z kal?r.
    if not required_exact.issubset(
        actual_exact
    ):
        return False

    # Y?zdeler i?in i? tolerans?:
    # required=99.999, actual=98.0 -> fark 1.999 -> kabul.
    # required=99.999, actual=96.0 -> fark 3.999 -> kabul edilmez.
    PERCENTAGE_POINT_TOLERANCE = 2.0

    for required in required_percentages:

        matched = any(
            (
                actual
                >= required
            )
            or (
                required - actual
                <= PERCENTAGE_POINT_TOLERANCE
            )
            for actual
            in actual_percentages
        )

        if not matched:
            return False

    return True


def _alignment_profile(
    requirement: str,
    company: str,
    semantic_score: float,
) -> dict[str, Any]:

    requirement_documents = (
        _document_ids(
            requirement
        )
    )

    company_documents = (
        _document_ids(
            company
        )
    )

    shared_documents = (
        requirement_documents
        & company_documents
    )

    requirement_features = (
        _technical_feature_terms(
            requirement
        )
    )

    company_features = (
        _technical_feature_terms(
            company
        )
    )

    shared_features = (
        requirement_features
        & company_features
    )
        # CBC, CBE, AMF vb. teknik varlık adları tek başına
    # bir yeteneğin karşılandığını kanıtlamaz.
    entity_only_features = {
        "cbc",
        "cbe",
        "amf",
        "mme",
        "bsc",
        "rnc",
        "enb",
        "gnb",
    }

    requirement_capability_features = (
        requirement_features
        - entity_only_features
    )

    shared_capability_features = (
        shared_features
        - entity_only_features
    )

    shared_entity_features = (
        shared_features
        & entity_only_features
    )

    requirement_keywords = (
        _keyword_tokens(
            requirement
        )
    )

    company_keywords = (
        _keyword_tokens(
            company
        )
    )

    shared_keywords = (
        requirement_keywords
        & company_keywords
    )

    feature_coverage = (
        len(shared_capability_features)
        / len(requirement_capability_features)
        if requirement_capability_features
        else 0.0
    )

    keyword_coverage = (
        len(shared_keywords)
        / len(requirement_keywords)
        if requirement_keywords
        else 0.0
    )

    requirement_values = (
        _critical_values(
            requirement
        )
    )

    company_values = (
        _critical_values(
            company
        )
    )

    critical_mismatch = (
        not _critical_values_satisfied(
            requirement_values,
            company_values,
        )
    )
    final_score = (
        semantic_score
        + (
            0.55
            if shared_documents
            else 0.0
        )
        + 0.35
        * feature_coverage
        + min(
            0.18,
            0.06
            * len(
                shared_capability_features
            ),
        )
        + min(
            0.04,
            0.02
            * len(
                shared_entity_features
            ),
        )
        + 0.10
        * keyword_coverage
    )

    strong = False

    if shared_documents:
        strong = True

    elif requirement_capability_features:

        if (
            feature_coverage >= 0.75
            and semantic_score >= 0.68
        ):
            strong = True

        elif (
            len(
                shared_capability_features
            ) >= 2
            and feature_coverage >= 0.50
            and semantic_score >= 0.70
        ):
            strong = True

    elif (
        not requirement_features
        and len(shared_keywords) >= 3
        and keyword_coverage >= 0.45
        and semantic_score >= 0.80
    ):
        strong = True

    medium = (
        not strong
        and (
            (
                bool(
                    shared_capability_features
                )
                and feature_coverage >= 0.34
                and semantic_score >= 0.68
            )
            or (
                len(shared_keywords) >= 2
                and keyword_coverage >= 0.28
                and semantic_score >= 0.80
            )
        )
    )

    return {
        "final_score": (
            final_score
        ),
        "strong": strong,
        "medium": medium,
        "shared_documents": (
            shared_documents
        ),
        "shared_features": (
            shared_features
        ),
        "shared_keywords": (
            shared_keywords
        ),
        "feature_coverage": (
            feature_coverage
        ),
        "keyword_coverage": (
            keyword_coverage
        ),
        "critical_mismatch": (
            critical_mismatch
        ),
    }


def _is_specification_requirement(
    row: ExtractedRow,
) -> bool:

    text = _normalized_match_text(
        row.text
    )

    if not text:
        return False

    if (
        "teknik gereksinim"
        in text
        and "madde" in text
        and "uygunluk" in text
    ):
        return False

    legend_patterns = (
        r"^fc\s*[|:-].*tam uygun",
        r"^pc\s*[|:-].*kismen uygun",
        r"^nc\s*[|:-].*uygun degil",
    )

    if any(
        re.search(
            pattern,
            text,
        )
        for pattern
        in legend_patterns
    ):
        return False

    return (
        len(text) >= 12
    )


def _is_company_capability_row(
    row: ExtractedRow,
) -> bool:
    values = [
        value
        for value
        in row.values.values()
        if _clean(value)
    ]

    if not values:
        return False

    if _header_score(values) >= 8:
        return False

    normalized_values = [
        _normalized_match_text(value)
        for value in values
        if value
    ]

    normalized_text = (
        _normalized_match_text(
            row.text
        )
    )

    status_literals = {
        "fully supported",
        "partly supported",
        "partially supported",
        "not supported",
        "full compliance",
        "partial compliance",
        "tam uygun",
        "kismen uygun",
        "uygun degil",
    }

    # Salt durum satiri capability degildir.
    if normalized_text in status_literals:
        return False

    section_labels = {
        "gui features",
        "gui",
        "integration",
        "dimensioning",
        "messages",
        "message features",
        "security",
        "authentication",
        "kimlik dogrulama",
        "authorization",
        "yetkilendirme",
        "logging",
        "loglama",
        "reporting",
        "raporlama",
        "database",
        "db",
        "network",
        "system",
        "general",
        "genel",
        "support",
        "operation",
        "operations",
        "maintenance",
        "bakim",
    }

    content_values = [
        value
        for value in normalized_values
        if value not in status_literals
    ]

    if (
        len(content_values) == 1
        and content_values[0]
        in section_labels
    ):
        return False

    # -----------------------------------------------------
    # FREE-TEXT SOURCES
    # -----------------------------------------------------
    #
    # XLSX/CSV eski tablo davranisini korur.
    # DOCX/PDF/OCR/PPTX/TXT gibi kaynaklarda ise
    # capability tek bir dogal cumle olabilir.
    #
    # "desteklenmektedir" -> capability
    # "desteklenmelidir"  -> requirement
    # -----------------------------------------------------

    free_text_kinds = {
        "docx",
        "pdf",
        "image",
        "pptx",
        "txt",
        "md",
        "json",
        "xml",
    }

    source_kind = (
        row.source_kind
        or ""
    ).strip().casefold()

    if source_kind in free_text_kinds:

        requirement_patterns = (
            r"\bdesteklenmelidir\b",
            r"\bdesteklenmeli\b",
            r"\bsaglanmalidir\b",
            r"\bsaglanmali\b",
            r"\bolmalidir\b",
            r"\bolmali\b",
            r"\bgereklidir\b",
            r"\bgerekir\b",
            r"\bzorunludur\b",
            r"\bshall\b",
            r"\bmust\b",
            r"\brequired\b",
            r"\bis required\b",
            r"\bshall support\b",
            r"\bmust support\b",
        )

        is_requirement = any(
            re.search(
                pattern,
                normalized_text,
                flags=re.IGNORECASE,
            )
            for pattern
            in requirement_patterns
        )

        if is_requirement:
            return False

        capability_patterns = (
            # Turkish positive
            r"\bdesteklenmektedir\b",
            r"\bdesteklenir\b",
            r"\bdesteklidir\b",
            r"\bdestegi bulunmaktadir\b",
            r"\bdestek bulunmaktadir\b",
            r"\bbulunmaktadir\b",
            r"\bmevcuttur\b",
            r"\bsaglanmaktadir\b",
            r"\bsaglanir\b",
            r"\bkullanilabilmektedir\b",
            r"\bkullanilabilir\b",

            # Turkish negative
            r"\bdesteklenmemektedir\b",
            r"\bdesteklenmez\b",
            r"\bbulunmamaktadir\b",
            r"\bmevcut degildir\b",
            r"\bsaglanmamaktadir\b",

            # English positive
            r"\bis supported\b",
            r"\bare supported\b",
            r"\bsupports\b",
            r"\bsupported\b",
            r"\bis available\b",
            r"\bare available\b",
            r"\bprovides\b",
            r"\bcapable of\b",
            r"\bhas support\b",

            # English negative
            r"\bnot supported\b",
            r"\bunsupported\b",
            r"\bnot available\b",
        )

        is_capability = any(
            re.search(
                pattern,
                normalized_text,
                flags=re.IGNORECASE,
            )
            for pattern
            in capability_patterns
        )

        if is_capability:
            return True

    # Kisa, tek hucreli ve teknik kanit
    # tasimayan bolum basliklarini at.
    if (
        len(normalized_values) == 1
        and _company_support_status(row)
        is None
        and not _document_ids(
            row.text
        )
        and len(
            normalized_text.split()
        ) <= 6
        and not re.search(
            r"[0-9?.:]",
            normalized_text,
        )
        and not _atomic_capability_terms(
            row.text
        )
    ):
        return False

    # Ayni basligin birden fazla kolonda
    # tekrarlandigi satirlari at.
    if (
        len(normalized_values) >= 2
        and len(
            set(normalized_values)
        ) == 1
        and _company_support_status(row)
        is None
    ):
        return False

    # Tek hucreli satirda durum, dokuman
    # veya teknik ozellik yoksa at.
    if (
        len(values) == 1
        and _company_support_status(row)
        is None
        and not _document_ids(
            row.text
        )
        and not _technical_feature_terms(
            row.text
        )
        and not _atomic_capability_terms(
            row.text
        )
    ):
        return False

    return (
        len(normalized_text)
        >= 8
    )




def _rows_from_xlsx(
    content: bytes,
) -> list[ExtractedRow]:

    workbook = load_workbook(
        io.BytesIO(content),
        read_only=False,
        data_only=True,
    )

    rows: list[ExtractedRow] = []

    for worksheet in workbook.worksheets:

        raw_rows = [
            [
                cell.value
                for cell in row
            ]
            for row
            in worksheet.iter_rows()
        ]

        if not raw_rows:
            continue

        # Propagate merged-cell values. This is important for
        # Compliance Status cells spanning multiple requirements.
        for merged_range in (
            worksheet.merged_cells.ranges
        ):

            value = worksheet.cell(
                merged_range.min_row,
                merged_range.min_col,
            ).value

            if value is None:
                continue

            for row_index in range(
                merged_range.min_row,
                merged_range.max_row + 1,
            ):

                for column_index in range(
                    merged_range.min_col,
                    merged_range.max_col + 1,
                ):

                    if (
                        row_index - 1
                        >= len(raw_rows)
                    ):
                        continue

                    row_values = (
                        raw_rows[
                            row_index - 1
                        ]
                    )

                    while (
                        len(row_values)
                        < column_index
                    ):
                        row_values.append(
                            None
                        )

                    if not _clean(
                        row_values[
                            column_index - 1
                        ]
                    ):
                        row_values[
                            column_index - 1
                        ] = value

        candidate_headers: list[
            tuple[int, int]
        ] = []

        for index, row in enumerate(
            raw_rows[:50]
        ):

            clean_values = [
                _clean(value)
                for value
                in row
            ]

            score = _header_score(
                clean_values
            )

            if score:
                candidate_headers.append(
                    (
                        score,
                        index,
                    )
                )

        if candidate_headers:

            candidate_headers.sort(
                reverse=True
            )

            header_index = (
                candidate_headers[
                    0
                ][1]
            )

        else:

            header_index = None

            for index, row in enumerate(
                raw_rows
            ):

                clean_values = [
                    _clean(value)
                    for value
                    in row
                ]

                if (
                    _meaningful(
                        clean_values
                    )
                    and sum(
                        bool(value)
                        for value
                        in clean_values
                    )
                    >= 2
                ):
                    header_index = index
                    break

        if header_index is None:
            continue

        header_values = [
            _clean(value)
            for value
            in raw_rows[
                header_index
            ]
        ]

        headers = _unique_headers(
            header_values
        )

        for row_index, row in enumerate(
            raw_rows[
                header_index + 1:
            ],
            start=header_index + 2,
        ):

            clean_values = [
                _clean(value)
                for value
                in row
            ]

            if not _meaningful(
                clean_values
            ):
                continue

            if _header_score(
                clean_values
            ) >= 8:
                continue

            values = {
                headers[index]:
                (
                    clean_values[index]
                    if index
                    < len(clean_values)
                    else ""
                )
                for index
                in range(
                    len(headers)
                )
            }

            useful = [
                value
                for value
                in values.values()
                if value
            ]

            text = " | ".join(
                useful
            )

            if len(text) < 8:
                continue

            rows.append(
                ExtractedRow(
                    row_number=(
                        row_index
                    ),
                    values=values,
                    text=text,
                )
            )

    return rows



def _rows_from_csv(
    content: bytes,
) -> list[ExtractedRow]:

    decoded = content.decode(
        "utf-8-sig",
        errors="replace",
    )

    reader = csv.DictReader(
        io.StringIO(decoded)
    )

    rows: list[ExtractedRow] = []

    for row_number, row in enumerate(
        reader,
        start=2,
    ):

        values = {
            _clean(key):
            _clean(value)
            for key, value
            in row.items()
            if key is not None
        }

        useful = [
            value
            for value
            in values.values()
            if value
        ]

        if not _meaningful(
            useful
        ):
            continue

        text = " | ".join(
            useful
        )

        if len(text) < 8:
            continue

        rows.append(
            ExtractedRow(
                row_number=row_number,
                values=values,
                text=text,
            )
        )

    return rows



def _decode_text_bytes(
    content: bytes,
) -> str:
    """
    TXT / MD / JSON / XML gibi metin dosyalarini
    olabildigince guvenli decode eder.
    """

    for encoding in (
        "utf-8-sig",
        "utf-8",
        "cp1254",
        "cp1252",
        "latin-1",
    ):
        try:
            return content.decode(
                encoding
            )
        except UnicodeDecodeError:
            continue

    raise ValueError(
        "Metin dosyasinin karakter kodlamasi "
        "cozumlenemedi."
    )


def _text_blocks(
    value: str,
) -> list[str]:
    """
    Serbest metni compliance motorunun kullanabilecegi
    mantikli parcalara ayirir.

    Once bos satirlarla paragraf ayirir.
    Cok uzun bloklari satir bazinda boler.
    """

    normalized = (
        value
        or ""
    ).replace(
        "\r\n",
        "\n",
    ).replace(
        "\r",
        "\n",
    )

    paragraphs = [
        re.sub(
            r"\s+",
            " ",
            part,
        ).strip()
        for part in re.split(
            r"\n\s*\n+",
            normalized,
        )
        if part.strip()
    ]

    blocks: list[str] = []

    for paragraph in paragraphs:

        if len(paragraph) <= 1200:
            blocks.append(
                paragraph
            )
            continue

        lines = [
            re.sub(
                r"\s+",
                " ",
                line,
            ).strip()
            for line in paragraph.split(
                "\n"
            )
            if line.strip()
        ]

        if lines:
            blocks.extend(
                lines
            )
        else:
            blocks.append(
                paragraph
            )

    return [
        block
        for block in blocks
        if _meaningful(
            [block]
        )
    ]


def _rows_from_plain_text(
    value: str,
    *,
    source_filename: str,
    source_kind: str,
    source_page: int | None = None,
    row_start: int = 1,
) -> list[ExtractedRow]:

    rows: list[ExtractedRow] = []

    for offset, block in enumerate(
        _text_blocks(
            value
        ),
    ):
        row_number = (
            row_start
            + offset
        )

        rows.append(
            ExtractedRow(
                row_number=row_number,
                values={
                    "text": block,
                },
                text=block,
                source_filename=(
                    source_filename
                ),
                source_page=(
                    source_page
                ),
                source_kind=(
                    source_kind
                ),
            )
        )

    return rows



# PARADOKS OCR HELPERS

_DEFAULT_TESSERACT_EXE = Path(
    r"C:\Program Files\Tesseract-OCR\tesseract.exe"
)

_DEFAULT_TESSDATA_PATH = Path(
    r"C:\paradoks_data_v3\tessdata"
)

_IMAGE_EXTENSIONS = {
    ".png",
    ".jpg",
    ".jpeg",
    ".webp",
    ".bmp",
    ".tif",
    ".tiff",
}


def _resolve_tesseract_executable(
) -> str:
    configured = (
        os.getenv(
            "PARADOKS_TESSERACT_CMD",
            "",
        ).strip()
    )

    if configured:
        candidate = Path(
            configured
        )

        if candidate.is_file():
            return str(
                candidate
            )

    discovered = shutil.which(
        "tesseract"
    )

    if discovered:
        return discovered

    if _DEFAULT_TESSERACT_EXE.is_file():
        return str(
            _DEFAULT_TESSERACT_EXE
        )

    raise ValueError(
        "Tesseract OCR bulunamadi. "
        "Tesseract kurulumunu kontrol edin."
    )


def _resolve_tessdata_path(
) -> Path | None:
    configured = (
        os.getenv(
            "PARADOKS_TESSDATA_PATH",
            "",
        ).strip()
    )

    if configured:
        candidate = Path(
            configured
        )

        if candidate.is_dir():
            return candidate

    if _DEFAULT_TESSDATA_PATH.is_dir():
        return _DEFAULT_TESSDATA_PATH

    executable = Path(
        _resolve_tesseract_executable()
    )

    bundled = (
        executable.parent
        / "tessdata"
    )

    if bundled.is_dir():
        return bundled

    return None


def _ocr_language(
) -> str:
    tessdata = (
        _resolve_tessdata_path()
    )

    if tessdata is None:
        return "eng"

    languages: list[str] = []

    if (
        tessdata
        / "tur.traineddata"
    ).is_file():
        languages.append(
            "tur"
        )

    if (
        tessdata
        / "eng.traineddata"
    ).is_file():
        languages.append(
            "eng"
        )

    return (
        "+".join(
            languages
        )
        or "eng"
    )


def _ocr_image_object(
    image: Image.Image,
) -> str:
    pytesseract.pytesseract.tesseract_cmd = (
        _resolve_tesseract_executable()
    )

    prepared = ImageOps.exif_transpose(
        image
    )

    if prepared.mode not in {
        "RGB",
        "L",
    }:
        prepared = prepared.convert(
            "RGB"
        )

    prepared = ImageOps.grayscale(
        prepared
    )

    prepared = ImageOps.autocontrast(
        prepared
    )

    # Cok kucuk ekran goruntulerinde OCR kalitesini
    # arttirmak icin makul bir olceklendirme uygula.
    if prepared.width < 1400:
        scale = min(
            2.0,
            1400
            / max(
                prepared.width,
                1,
            ),
        )

        if scale > 1.15:
            prepared = prepared.resize(
                (
                    int(
                        prepared.width
                        * scale
                    ),
                    int(
                        prepared.height
                        * scale
                    ),
                )
            )

    tessdata = (
        _resolve_tessdata_path()
    )

    config_parts = [
        "--psm 6",
    ]

    # Windows'ta pytesseract'in --tessdata-dir
    # argumanina verilen quoted path bazen
    # yanlis parse ediliyor.
    #
    # TESSDATA_PREFIX daha stabil.
    if tessdata is not None:
        os.environ[
            "TESSDATA_PREFIX"
        ] = str(
            tessdata
        )

    try:
        value = (
            pytesseract.image_to_string(
                prepared,
                lang=_ocr_language(),
                config=" ".join(
                    config_parts
                ),
            )
        )
    except Exception as error:
        raise ValueError(
            "OCR islemi basarisiz oldu."
        ) from error

    return re.sub(
        r"[ \t]+",
        " ",
        value or "",
    ).strip()


def _rows_from_image(
    content: bytes,
    *,
    source_filename: str,
) -> list[ExtractedRow]:
    try:
        image = Image.open(
            io.BytesIO(
                content
            )
        )
    except Exception as error:
        raise ValueError(
            (
                "Gorsel dosyasi okunamadi: "
                f"{source_filename}"
            )
        ) from error

    rows: list[
        ExtractedRow
    ] = []

    frame_count = int(
        getattr(
            image,
            "n_frames",
            1,
        )
        or 1
    )

    next_row = 1

    for frame_index in range(
        frame_count
    ):
        try:
            image.seek(
                frame_index
            )
        except EOFError:
            break

        frame = image.copy()

        text_value = (
            _ocr_image_object(
                frame
            )
        )

        frame_rows = (
            _rows_from_plain_text(
                text_value,
                source_filename=(
                    source_filename
                ),
                source_kind="image",
                source_page=(
                    frame_index + 1
                ),
                row_start=(
                    next_row
                ),
            )
        )

        rows.extend(
            frame_rows
        )

        next_row += len(
            frame_rows
        )

    if not rows:
        raise ValueError(
            (
                "Gorselden anlamli metin "
                f"cikarilamadi: "
                f"{source_filename}"
            )
        )

    return rows


def _rows_from_xls(
    content: bytes,
    *,
    source_filename: str,
) -> list[ExtractedRow]:
    try:
        workbook = (
            xlrd.open_workbook(
                file_contents=content
            )
        )
    except Exception as error:
        raise ValueError(
            (
                "XLS dosyasi okunamadi: "
                f"{source_filename}"
            )
        ) from error

    rows: list[
        ExtractedRow
    ] = []

    row_number = 1

    for sheet in workbook.sheets():

        for source_row_index in range(
            sheet.nrows
        ):
            values = [
                _clean(
                    sheet.cell_value(
                        source_row_index,
                        column_index,
                    )
                )
                for column_index
                in range(
                    sheet.ncols
                )
            ]

            meaningful_values = [
                value
                for value in values
                if value
            ]

            if not _meaningful(
                meaningful_values
            ):
                continue

            row_text = " | ".join(
                meaningful_values
            )

            rows.append(
                ExtractedRow(
                    row_number=(
                        row_number
                    ),
                    values={
                        "sheet": (
                            sheet.name
                        ),
                        **{
                            (
                                f"column_"
                                f"{index + 1}"
                            ): value
                            for index, value
                            in enumerate(
                                meaningful_values
                            )
                        },
                    },
                    text=row_text,
                    source_filename=(
                        source_filename
                    ),
                    source_kind="xls",
                )
            )

            row_number += 1

    if not rows:
        raise ValueError(
            (
                "XLS dosyasindan anlamli "
                f"icerik cikarilamadi: "
                f"{source_filename}"
            )
        )

    return rows


def _rows_from_pdf(
    content: bytes,
    *,
    source_filename: str,
) -> list[ExtractedRow]:
    try:
        reader = PdfReader(
            io.BytesIO(
                content
            )
        )
    except Exception as error:
        raise ValueError(
            (
                "PDF okunamadi: "
                f"{source_filename}"
            )
        ) from error

    rows: list[
        ExtractedRow
    ] = []

    next_row_number = 1

    mupdf_document = None

    try:
        for page_index, page in enumerate(
            reader.pages
        ):
            page_number = (
                page_index + 1
            )

            try:
                page_text = (
                    page.extract_text()
                    or ""
                )
            except Exception:
                page_text = ""

            compact_text = re.sub(
                r"\s+",
                "",
                page_text,
            )

            # Metin katmani yoksa veya anlamsiz derecede
            # azsa sayfayi goruntuye cevirip OCR uygula.
            if len(compact_text) < 40:

                if mupdf_document is None:
                    try:
                        mupdf_document = (
                            pymupdf.open(
                                stream=content,
                                filetype="pdf",
                            )
                        )
                    except Exception as error:
                        raise ValueError(
                            (
                                "PDF OCR icin "
                                "goruntulenemedi: "
                                f"{source_filename}"
                            )
                        ) from error

                try:
                    pdf_page = (
                        mupdf_document.load_page(
                            page_index
                        )
                    )

                    pixmap = (
                        pdf_page.get_pixmap(
                            matrix=(
                                pymupdf.Matrix(
                                    2.5,
                                    2.5,
                                )
                            ),
                            alpha=False,
                        )
                    )

                    image = Image.open(
                        io.BytesIO(
                            pixmap.tobytes(
                                "png"
                            )
                        )
                    )

                    ocr_text = (
                        _ocr_image_object(
                            image
                        )
                    )

                    if len(
                        re.sub(
                            r"\s+",
                            "",
                            ocr_text,
                        )
                    ) > len(
                        compact_text
                    ):
                        page_text = (
                            ocr_text
                        )

                except Exception as error:
                    if not compact_text:
                        raise ValueError(
                            (
                                "PDF sayfasi OCR ile "
                                "okunamadi: "
                                f"{source_filename}, "
                                f"sayfa {page_number}"
                            )
                        ) from error

            page_rows = (
                _rows_from_plain_text(
                    page_text,
                    source_filename=(
                        source_filename
                    ),
                    source_kind="pdf",
                    source_page=(
                        page_number
                    ),
                    row_start=(
                        next_row_number
                    ),
                )
            )

            rows.extend(
                page_rows
            )

            next_row_number += len(
                page_rows
            )

    finally:
        if mupdf_document is not None:
            mupdf_document.close()

    if not rows:
        raise ValueError(
            (
                "PDF dosyasindan anlamli "
                f"icerik cikarilamadi: "
                f"{source_filename}"
            )
        )

    return rows


def _rows_from_docx(
    content: bytes,
    *,
    source_filename: str,
) -> list[ExtractedRow]:

    try:
        document = Document(
            io.BytesIO(
                content
            )
        )
    except Exception as error:
        raise ValueError(
            f"DOCX okunamadi: "
            f"{source_filename}"
        ) from error

    rows: list[ExtractedRow] = []
    row_number = 1

    for paragraph in document.paragraphs:

        clean_text = _clean(
            paragraph.text
        )

        if not _meaningful(
            [clean_text]
        ):
            continue

        rows.append(
            ExtractedRow(
                row_number=row_number,
                values={
                    "text": clean_text,
                },
                text=clean_text,
                source_filename=(
                    source_filename
                ),
                source_kind="docx",
            )
        )

        row_number += 1

    # DOCX icindeki tablolar da teknik sartnamelerde
    # cok kullaniliyor; bunlari atlamiyoruz.
    for table in document.tables:

        for table_row in table.rows:

            values = [
                _clean(
                    cell.text
                )
                for cell in (
                    table_row.cells
                )
            ]

            values = [
                value
                for value in values
                if value
            ]

            if not _meaningful(
                values
            ):
                continue

            joined = " | ".join(
                values
            )

            rows.append(
                ExtractedRow(
                    row_number=(
                        row_number
                    ),
                    values={
                        f"column_{index + 1}": value
                        for index, value
                        in enumerate(
                            values
                        )
                    },
                    text=joined,
                    source_filename=(
                        source_filename
                    ),
                    source_kind="docx",
                )
            )

            row_number += 1

    if not rows:
        raise ValueError(
            (
                "DOCX dosyasindan anlamli "
                f"icerik cikarilamadi: "
                f"{source_filename}"
            )
        )

    return rows


def _rows_from_pptx(
    content: bytes,
    *,
    source_filename: str,
) -> list[ExtractedRow]:

    try:
        presentation = Presentation(
            io.BytesIO(
                content
            )
        )
    except Exception as error:
        raise ValueError(
            f"PPTX okunamadi: "
            f"{source_filename}"
        ) from error

    rows: list[ExtractedRow] = []
    row_number = 1

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1,
    ):

        slide_parts: list[str] = []

        for shape in slide.shapes:

            if not hasattr(
                shape,
                "text",
            ):
                continue

            value = _clean(
                getattr(
                    shape,
                    "text",
                    "",
                )
            )

            if value:
                slide_parts.append(
                    value
                )

        slide_text = " | ".join(
            slide_parts
        )

        if not _meaningful(
            [slide_text]
        ):
            continue

        rows.append(
            ExtractedRow(
                row_number=row_number,
                values={
                    "text": slide_text,
                },
                text=slide_text,
                source_filename=(
                    source_filename
                ),
                source_page=(
                    slide_number
                ),
                source_kind="pptx",
            )
        )

        row_number += 1

    if not rows:
        raise ValueError(
            (
                "PPTX dosyasindan anlamli "
                f"icerik cikarilamadi: "
                f"{source_filename}"
            )
        )

    return rows


def _rows_from_json(
    content: bytes,
    *,
    source_filename: str,
) -> list[ExtractedRow]:

    raw_text = _decode_text_bytes(
        content
    )

    try:
        parsed = json.loads(
            raw_text
        )
    except json.JSONDecodeError as error:
        raise ValueError(
            f"JSON gecersiz: "
            f"{source_filename}"
        ) from error

    formatted = json.dumps(
        parsed,
        ensure_ascii=False,
        indent=2,
    )

    return _rows_from_plain_text(
        formatted,
        source_filename=(
            source_filename
        ),
        source_kind="json",
    )


def _annotate_rows(
    rows: list[ExtractedRow],
    *,
    source_filename: str,
    source_kind: str,
) -> list[ExtractedRow]:
    """
    Mevcut XLSX/CSV parserlarini degistirmeden
    kaynak dosya bilgisini ekler.
    """

    for row in rows:

        if not row.source_filename:
            row.source_filename = (
                source_filename
            )

        if not row.source_kind:
            row.source_kind = (
                source_kind
            )

    return rows


def extract_rows(
    *,
    filename: str,
    content: bytes,
) -> list[ExtractedRow]:

    if not content:
        raise ValueError(
            f"Dosya bos: {filename}"
        )

    suffix = Path(
        filename
    ).suffix.casefold()

    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            (
                "Desteklenmeyen dosya "
                f"formati: "
                f"{suffix or '(uzantisiz)'}. "
                f"Dosya: {filename}"
            )
        )

    if suffix == ".xlsx":
        return _annotate_rows(
            _rows_from_xlsx(
                content
            ),
            source_filename=(
                filename
            ),
            source_kind="xlsx",
        )

    if suffix == ".xls":
        return _rows_from_xls(
            content,
            source_filename=(
                filename
            ),
        )

    if suffix == ".csv":
        return _annotate_rows(
            _rows_from_csv(
                content
            ),
            source_filename=(
                filename
            ),
            source_kind="csv",
        )

    if suffix == ".pdf":
        return _rows_from_pdf(
            content,
            source_filename=(
                filename
            ),
        )

    if suffix == ".docx":
        return _rows_from_docx(
            content,
            source_filename=(
                filename
            ),
        )

    if suffix == ".pptx":
        return _rows_from_pptx(
            content,
            source_filename=(
                filename
            ),
        )

    if suffix == ".json":
        return _rows_from_json(
            content,
            source_filename=(
                filename
            ),
        )

    if suffix in {
        ".txt",
        ".md",
        ".xml",
    }:
        decoded = (
            _decode_text_bytes(
                content
            )
        )

        rows = (
            _rows_from_plain_text(
                decoded,
                source_filename=(
                    filename
                ),
                source_kind=(
                    suffix.lstrip(".")
                ),
            )
        )

        if not rows:
            raise ValueError(
                (
                    "Dosyadan anlamli metin "
                    f"cikarilamadi: "
                    f"{filename}"
                )
            )

        return rows

    if suffix in _IMAGE_EXTENSIONS:
        return _rows_from_image(
            content,
            source_filename=(
                filename
            ),
        )

    raise ValueError(
        (
            "Dosya formati "
            f"islenemedi: {filename}"
        )
    )


def _cosine_scores(
    matrix,
    vector,
):
    # EmbeddingService normalize_embeddings=True
    # kulland??? i?in dot product = cosine similarity.
    return matrix @ vector


def _status_from_score(
    score: float,
) -> str:

    # Bunlar MVP g?ven e?ikleridir.
    # Belirsiz e?le?meler otomatik NC yap?lmaz.
    if score >= 0.83:
        return "FC"

    if score >= 0.72:
        return "PC"

    if score >= 0.58:
        return "REVIEW"

    return "NC"

def _requirement_body(
    row: ExtractedRow,
) -> str:
    """Şartname satırındaki gerçek teknik gereksinim hücresini döndürür."""
    for key, value in row.values.items():
        normalized_key = _normalized_match_text(key)
        if (
            "teknik gereksinim" in normalized_key
            or normalized_key in {
                "requirement",
                "technical requirement",
                "technical requirements",
            }
        ):
            cleaned = _clean(value)
            if cleaned:
                return cleaned

    # CSV / beklenmeyen başlık yapılarında güvenli fallback.
    return row.text

def _atomic_capability_terms(
    value: str,
) -> set[str]:
    """Compliance kararı için metni atomik teknik zorunluluklara ayırır."""
    text = _normalized_match_text(value)
    found: set[str] = set()

    patterns = (
        # PARADOKS CORE COMPLIANCE CAPABILITIES
        (
            "ipv6",
            r"\bipv6\b",
        ),
        (
            "http2",
            r"\bhttp\W*2\b|\bhttp2\b",
        ),
        (
            "active_active",
            r"\baktif\W*aktif\b",
        ),
        (
            "availability",
            r"\bkullanilabilirli(?:k|g)\w*\b",
        ),
        (
            "aes_256",
            r"\baes\W*256\b",
        ),
        (
            "two_factor_authentication",
            (
                r"\b2fa\b"
                r"|\btwo[-\s]+factor authentication\b"
                r"|\bcift faktor\w* kimlik dogrulama\b"
                r"|\biki faktor\w* kimlik dogrulama\b"
            ),
        ),

        # Mobil erişim teknolojileri
        ("tech_cdma", r"\bcdma\b"),
        ("tech_gsm", r"\bgsm\b|\b2g\b"),
        ("tech_umts", r"\bumts\b|\b3g\b"),
        ("tech_lte", r"\blte\b|\b4g\b|\b4[.]5g\b"),
        ("tech_5g", r"\b5g(?: nr)?\b"),
        ("cell_broadcast_service", r"\bcell broadcast service\b|\bhucre yayin hizmet\w*\b|\bcbs\b"),

        # Sanallaştırma / donanım / HA
        ("vm", r"\bvm\b|\bvirtual machine\b|\bsanal makine\b"),
        ("vm_sizing", r"\bvm sizing\b|\bvm boyutlandir\w*\b|\bvirtual machine sizing\b|\bsize of requirements\b"),
        ("testbed", r"\btestbed\b|\btest system\b|\btest ortam\w*\b"),
        ("geo_site", r"\bgeo-site\b|\bgeo site\b|\bcografi saha\b|\b2 location\b|\btwo locations\b|\biki lokasyon\b"),
        ("x86", r"\bx86\b"),
        ("vmware", r"\bvmware\b"),
        ("nutanix", r"\bnutanix\b"),
        ("active_active", r"\bactive-active\b"),
        ("active_standby", r"\bactive-standby\b"),
        ("geo_redundancy", r"\bgeo-redundan\w*\b|\bcografi yedeklilik\b"),
        ("availability", r"\bavailability\b|\berisilebilir\w*\b"),
        ("scalability", r"\bscalab\w*\b|\bolceklen\w*\b"),
        ("replication", r"\breplication\w*\b|\breplikasyon\w*\b"),
        ("backup", r"\bbackup\b|\byedekle(?:me|mek|nir|nmesi|nmelidir)?\b"),
        ("restore", r"\brestore\b|\bgeri yukle\w*\b"),
        ("failover", r"\bfailover\b|\byuk devret\w*\b"),
        ("interface_redundancy", r"\binterface redundan\w*\b|\barabirim yedekliligi\b"),

        # GUI / harita
        ("gui", r"\bgui\b|\bgrafik kullanici arabirimi\b|\bweb gui\b"),
        ("map_gis", r"\bgis\b|\bmap integration\b|\bharita entegrasyon\w*\b"),
        ("gis_license", r"\b(?:map|gis|harita)[^|.;]*lisans\w*\b|\blisans\w*[^|.;]*(?:map|gis|harita)\b"),
        (
            "warning_source_integration",
            r"\b(?:multiple|birden fazla)[^|.;]*(?:warning source|uyari kaynag)\w*\b"
            r"|\bintegration with other public alerting systems\b"
            r"|\bdiger kamu uyari sistem\w*[^|.;]*entegrasyon\w*\b",
        ),
        (
            "region_targeting",
            r"\bbolgelere gore hedefle\w*\b"
            r"|\bprovince\b|\bdistrict\b"
            r"|\bselected map area\b"
            r"|\bil\b[^|.;]*\bilce\b",
        ),
        (
            "target_cell",
            r"\bcell based targeting\b"
            r"|\bcell selection\b"
            r"|\bcell db\b"
            r"|\bcell\b[^|.;]{0,45}\bselection\b"
            r"|\b(?:province|district|selected map area)[^|.;]{0,80}\bcell\b"
            r"|\bhucre[^|.;]*hedefle\w*\b",
        ),
        ("target_sector", r"\bsector[^|.;]*target\w*\b|\bsektor[^|.;]*hedefle\w*\b"),
        ("target_cgi", r"\bcgi\b"),

        # Mesaj fonksiyonları
        ("message_schedule", r"\bmessage schedul\w*\b|\bmesaj zamanla\w*\b|\bschedule function\b"),
        ("message_repeat", r"\brepeat function\b|\brepeated message\b|\bmessage repeat\w*\b|\btekrarli mesaj\b|\btekrarli mesaj yayin\w*\b"),
        (
            "message_segmentation",
            r"\bsegmentation\b"
            r"|\bmessage segment\w*\b"
            r"|\bsegmented message\w*\b"
            r"|\bmesaj bolumle\w*\b"
            r"|\buzun mesaj\w*[^|.;]*bolumle\w*\b",
        ),
        ("message_priority", r"\bpriority for alerts\b|\bmessage priorit\w*\b|\bmesaj oncelik\w*\b"),
        ("overload_protection", r"\boverload protection\b|\basiri yuk koruma\w*\b"),
        ("message_update", r"\bmessage update\b|\bmesaj guncelle\w*\b|\baktif uyari mesaj\w*[^|.;]*guncellen\w*\b|\benhanced message update\b"),
        ("message_cancel", r"\bmessage cancellation\b|\bmessage cancel\w*\b|\bmesaj iptal\w*\b|\baktif uyari mesaj\w*[^|.;]*iptal\w*\b|\bcancellation features\b"),
        ("message_expiration", r"\bmessage expiration\b|\bexpiration\b|\bmesaj gecerlilik sonu\b|\bgecerlilik sonu\b"),
        ("message_statistics", r"\bmessage processing statistics\b|\bmessage status\w*[^|.;]*statistic\w*\b|\bmesaj isleme istatistik\w*\b|\bmesaj durum\w*[^|.;]*istatistik\w*\b"),
        ("concurrent_campaigns", r"\bconcurrent[^|.;]*(?:broadcast|campaign|session)\w*\b|\beszamanli[^|.;]*(?:yayin|kampanya|oturum)\w*\b"),

        # Protokol / prosedür
        ("cap", r"\bcap\b|\bcommon alerting protocol\b|\bortak uyari protokolu\b"),
        ("cap_v1_2", r"\bcap\s*v?1[.]2\b"),
        ("cbsp", r"\bcbsp\b"),
        ("cbsp_write_replace", r"\bcbsp[^|.;]*write-replace\b|\bwrite-replace[^|.;]*cbsp\b"),
        ("cbsp_kill", r"\bcbsp[^|.;]*kill\b|\bkill[^|.;]*cbsp\b"),
        ("s1ap_write_replace", r"\bs1ap[^|.;]*write-replace\b|\bwrite-replace[^|.;]*s1ap\b"),
        ("s1ap_kill", r"\bs1ap[^|.;]*kill\b|\bkill[^|.;]*s1ap\b"),
        ("sabp", r"\bsabp\b|\biu-?bc\b"),
        ("sbc", r"\bsbc\b"),
        ("n50", r"\bn50\b"),
        (
            "amf_integration",
            r"\bamf\b[^|.;]*(?:integration|entegrasyon|interface|arabirim)\w*"
            r"|\b(?:integration|entegrasyon|interface|arabirim)\w*[^|.;]*\bamf\b",
        ),
        ("pws", r"\bpws\b|\bpublic warning system\b|\bkamu uyari sistemi\b"),
        ("ack_nack", r"\back/nack\b|\back\b[^|.;]*\bnack\b"),
        ("xml_schema", r"\bxml schema\b|\bxml sema\b"),

        # Raporlama / dışa aktarım
        ("report_schedule", r"\bscheduled report\b|\breport schedul\w*\b|\bzamanlanmis rapor\b|\brapor zamanla\w*\b"),
        ("report_export", r"\breport[^|.;]*\bexport\b|\bexport[^|.;]*\breport\b|\brapor[^|.;]*disa aktar\w*\b|\bdisa aktar\w*[^|.;]*rapor\b"),
        ("email_delivery", r"\bemail\b|\be-mail\b|\be-posta\b"),
        ("sftp", r"\bsftp\b"),
        ("external_export", r"\bexternal platform\w*[^|.;]*export\b|\bexport[^|.;]*external platform\w*\b|\bharici platform\w*[^|.;]*disa aktar\w*\b|\bdisa aktar\w*[^|.;]*harici platform\w*\b"),

        # Kimlik / güvenlik
        ("rbac", r"\brbac\b|\brole-based access control\b|\brol-tabanli erisim kontrol\w*\b|\brol-yetki\b|\brol yetki matrisi\b|\byetki matrisi\b"),
        ("account_lockout", r"\baccount lockout\b|\baccount lock\w*\b|\bhesap kilit\w*\b|\bkullanici hesab\w*[^|.;]*kilit\w*\b|\bhatali giris\w*[^|.;]*kilit\w*\b"),
        ("tls", r"\btls\b"),
        ("tls_12", r"\btls\s*v?1[.]2\b"),
        ("tls_13", r"\btls\s*v?1[.]3\b"),
        ("https", r"\bhttps\b"),
        ("mutual_tls", r"\bmutual tls\b|\bmtls\b|\bkarsilikli tls\b"),
        ("syslog", r"\bsyslog\b"),
        (
            "siem_integration",
            r"\bsiem\b[^|.;]*(?:integration|integrable|compliance|uyum|entegrasyon)\w*"
            r"|\b(?:integration|integrable|compliance|uyum|entegrasyon)\w*[^|.;]*\bsiem\b"
            r"|\bsiem system\b"
            r"|\bsiem\b",
        ),
        ("snmpv3", r"\bsnmp\s*v?3\b|\bsnmpv3\b"),
        ("rest_api", r"\brest(?:ful)?(?: northbound)? api\b|\brest\b[^|.;]{0,40}\bapi\b"),
        ("ldap", r"\bldap\b|\bldaps\b"),
        (
            "active_directory",
            r"\bactive directory\b"
            r"|\bad\s*/\s*ldap\b"
            r"|\bldap\s*/\s*ad\b"
            r"|\bad\s+(?:and|or|ve|veya)\s+ldap\b",
        ),
        ("umbrella_integration", r"\bumbrella\b[^|.;]*(?:integration|entegrasyon)\w*\b|\b(?:integration|entegrasyon)\w*[^|.;]*umbrella\b"),
        ("certificate_lifecycle", r"\bcertificate lifecycle\b|\bcertificate[^|.;]*(?:renew\w*|replacement|expiry|expiration|install\w*)\b|\bsertifika yasam dongusu\b|\bsertifika[^|.;]*(?:kurulum\w*|yenile\w*|degistir\w*|sure sonu)\b"),
        ("certificate_authentication", r"\bclient ssl certificate authentication\b|\bcertificate authentication\b|\bsertifika ile kimlik dogrulama\b"),
        ("encryption_at_rest", r"\bencryption at rest\b|\bat-rest encryption\b|\bat rest\b[^|.;]*encrypt\w*\b|\bbekleyen[^|.;]*hassas veri\w*[^|.;]*sifrelen\w*\b|\bhassas veri\w*[^|.;]*bekleyen[^|.;]*sifrelen\w*\b"),
        ("audit_log", r"\baudit(?: level)? logs?\b|\baudit log\b|\bdenetim gunluk\w*\b"),
        ("digital_signature", r"\bdigital signature\b|\bdijital imza\b"),

        # Diğer
        ("support_24x7", r"\b7/24\b|\b24 hour\w*[^|.;]*7 day\w*\b"),
        (
            "message_delivery_delay",
            r"\bmessage sending delay\b"
            r"|\bmesagge sending delay\b"
            r"|\bcell broadcast[^|.;]*iletim suresi\w*\b"
            r"|\bhucre yayin(?:i)?[^|.;]*iletim suresi\w*\b"
            r"|\bmesaj[^|.;]*iletim suresi\w*\b",
        ),
        ("time_sync", r"\btime synchroni[sz]\w*\b|\bzaman senkron\w*\b|\bntp\b"),
        ("unicode", r"\bunicode\b|\bucs-?2\b"),
        ("gsm7", r"\bgsm\s*7-?bit\b"),
        ("ecgi", r"\becgi\b"),
        (
            "tracking_area",
            r"\btracking area\b"
            r"|\bizleme alani\b"
            r"|\bta\b[^|.;]{0,80}\b(?:target\w*|hedefle\w*)\b"
            r"|\b(?:target\w*|hedefle\w*)[^|.;]{0,80}\bta\b",
        ),
        ("nr_cell_id", r"\bnr cell id\b|\bnr hucre kimligi\b"),
        ("dbgf", r"\bdbgf\b|\bdevice-based geo\w*\b"),
        ("multi_vendor", r"\bmulti-vendor\b|\bcok ureticili\b"),
        ("multi_language", r"\bmulti-language\b|\bcok dilli\b|\bwhich languages supported\b|\blanguages? support\w*\b"),
        ("lang_uzbek", r"\buzbek\w*\b|\bozbek\w*\b"),
        ("lang_russian", r"\brussian\b|\brusca\b"),
        ("lang_english", r"\benglish\b|\bingilizce\b"),
        ("lang_karakalpak", r"\bkarakalpak\w*\b"),
    )

    for name, pattern in patterns:
        if re.search(pattern, text, flags=re.IGNORECASE):
            found.add(name)

    # Parola politikası cümleleri çoğu zaman özelliği isim vermeden listeler.
    password_context = (
        "parola" in text
        or "password" in text
        or "sifre" in text
    )
    if password_context:
        if re.search(r"karmasik\w*|complexity|buyuk harf[^|.;]*kucuk harf|kucuk harf[^|.;]*buyuk harf", text):
            found.add("password_complexity")
        if re.search(r"minimum uzunluk|minimum \d+\s*(?:karakter|character)|password[^|.;]*minimum \d+|parola\w*[^|.;]*minimum \d+", text):
            found.add("password_min_length")
        if re.search(r"gecerlilik suresi|password (?:expiry|validity|expiration)|parola\w*[^|.;]*(?:sure|gun)", text):
            found.add("password_expiry")
        if re.search(r"parola gecmisi|password history|son (?:bes|5) parola|previous password", text):
            found.add("password_history")

    # "AD entegrasyonu mümkün değilse..." bir AD desteği kanıtı değildir.
    if (
        "active directory ile entegrasyon mumkun degilse"
        in text
        or "if active directory integration is not possible"
        in text
    ):
        found.discard("active_directory")
        found.discard("ldap")

    # Eğitim satırında CBS kelimesi geçmesi, teknik CBS desteği kanıtı değildir.
    if "training" in text or "egitim" in text:
        found.discard("cell_broadcast_service")

    # LDAP/AD satırlarından ortak merkezi dizin entegrasyonu sinyali üret.
    if "ldap" in found or "active_directory" in found:
        found.add("central_directory")
    else:
        found.discard("central_directory")

    # At-rest encryption maddesindeki 'yedekler' kelimesi backup yeteneği değildir.
    if "encryption_at_rest" in found:
        found.discard("backup")


    # --------------------------------------------------------
    # AUDIT-SAFE ATOMIC POST-PROCESS
    # --------------------------------------------------------

    # "Uzbekistan" ?lke ad?d?r; tek ba??na ?zbek?e dil
    # deste?i anlam?na gelmez.
    if (
        "lang_uzbek" in found
        and "uzbekistan" in text
        and not re.search(
            r"\b(?:uzbek language|uzbek til|uzbekce|ozbekce|"
            r"language uzbek|uzbek language support)\b",
            text,
        )
    ):
        found.discard(
            "lang_uzbek"
        )

    # Mesaj?n effective/expiration zaman?n? destekleyen sat?rlar,
    # "message expiration" yetene?i i?in ge?erli kan?tt?r.
    if (
        re.search(
            r"\b(?:message[^|.;]*)?"
            r"(?:expiration|expirations|expiry)\b",
            text,
        )
        or re.search(
            r"\beffective times?[^|.;]*"
            r"(?:expiration|expirations|expiry)\b",
            text,
        )
    ):
        found.add(
            "message_expiration"
        )

    return found





def _requirement_kind(
    value: str,
) -> str:
    """
    Şartname maddesinin temel karar türünü belirler.

    CAPABILITY  : Üründe/yetenekte doğrudan var-yok kontrolü.
    INFO        : Tedarikçiden açıklama, liste, değer veya doküman istenir.
    COMMERCIAL  : Lisans, tedarik süresi, ticari destek gibi teklif teyidi gerekir.
    """
    text = _normalized_match_text(value)

    commercial_patterns = (
        r"\blisans sinirlam\w*\b",
        r"\ben az \d+\s*yil[^|.;]*(?:teknik destek|tedarik)\b",
        r"\blicen[cs]e limitation\b",
        r"\bcommercial\b",
    )
    if any(re.search(pattern, text) for pattern in commercial_patterns):
        return "COMMERCIAL"

    info_patterns = (
        r"\blisteleyiniz\b",
        r"\blistelemelidir\b",
        r"\blistelenmel\w*\b",
        r"\bliste sunmalidir\b",
        r"\baciklama sunmalidir\b",
        r"\baciklamalidir\b",
        r"\baciklamasini saglamalidir\b",
        r"\breferans listesi sunmalidir\b",
        r"\bmesaj formatini saglamalidir\b",
        r"\bparametreleri listelemelidir\b",
        r"\bmaksimum[^|.;]*belirtmelidir\b",
        r"\biletim suresini belirtmelidir\b",
        r"\bmesaj uzunlugunu belirtmelidir\b",
        r"\bdesteklenen[^|.;]*(?:saglayici|platform|standart|sartname)[^|.;]*belirtmelidir\b",
        r"\bprovide (?:a |an )?(?:description|list|reference list|message format)\b",
        r"\blist (?:all|the)\b",
        r"\bspecify (?:the )?(?:maximum|message|supported)\b",
    )

    if any(re.search(pattern, text) for pattern in info_patterns):
        return "INFO"

    # "Tedarikçi/Üretici ... belirtmelidir/açıklamalıdır" yapısı,
    # teknik bir özellikten ziyade teklif cevabı/dokümantasyon talebidir.
    if (
        any(subject in text for subject in ("tedarikci", "uretici", "vendor"))
        and re.search(
            r"\b(?:belirtmelidir|aciklamalidir|sunmalidir|listelemelidir|saglamalidir)\b",
            text,
        )
        and not re.search(
            r"\bdesteklenmelidir\b|\bdesteklemelidir\b|\bsupport\w*\b",
            text,
        )
    ):
        return "INFO"

    return "CAPABILITY"


def _information_requires_explicit_value(
    value: str,
) -> bool:
    text = _normalized_match_text(value)

    patterns = (
        r"\bmaksimum\b",
        r"\bmaximum\b",
        r"\biletim suresini belirt\w*\b",
        r"\bmessage delivery time\b",
        r"\bmessage length\b",
        r"\bmesaj uzunlugunu belirt\w*\b",
        r"\bsayisini belirt\w*\b",
    )

    return any(
        re.search(pattern, text)
        for pattern in patterns
    )


def _evidence_has_explicit_value(
    evidence_items: list[dict[str, Any]],
) -> bool:
    """
    5.7.1 gibi madde numaralarını gerçek teknik değer saymadan,
    ölçülebilir bir cevap bulunup bulunmadığını kontrol eder.
    """
    for item in evidence_items:
        text = _normalized_match_text(
            str(item.get("text", ""))
        )

        # 99.99%, 30 days, 1000 cells, 5 million vb.
        if re.search(
            r"\b\d+(?:[.,]\d+)?\s*(?:%|ms|s|sec|second|seconds|sn|dk|min|hour|hours|saat|day|days|gun|million|milyon|cell|cells|hucre|connection|connections|baglanti)\b",
            text,
        ):
            return True

        # Saf sayı/değer sütunu: "... | 65535 | Fully Supported" gibi.
        parts = [
            part.strip()
            for part in text.split("|")
            if part.strip()
        ]
        for part in parts:
            if re.fullmatch(
                r"\d+(?:[.,]\d+)?",
                part,
            ):
                return True

    return False

def _aggregate_requirement_evidence(
    *,
    requirement: ExtractedRow,
    company_rows: list[ExtractedRow],
    semantic_scores,
    scored_candidates: list[
        tuple[
            float,
            int,
            dict[str, Any],
        ]
    ],
) -> dict[str, Any] | None:
    """Bir şartname maddesini birden fazla şirket satırıyla birlikte değerlendirir."""
    requirement_text = _requirement_body(requirement)

    requirement_documents = _document_ids(requirement_text)
    requirement_capabilities = _atomic_capability_terms(requirement_text)

    # LDAP ve/veya AD gibi alternatifli gereksinimlerde iki teknolojiyi
    # aynı anda zorunlu tutma; ortak merkezi dizin atomu yeterlidir.
    normalized_requirement = _normalized_match_text(requirement_text)
    if (
        {"ldap", "active_directory"}.issubset(requirement_capabilities)
        and (
            "ve/veya" in normalized_requirement
            or "and/or" in normalized_requirement
            or "ldap veya" in normalized_requirement
            or "ldap ve/veya" in normalized_requirement
        )
    ):
        requirement_capabilities.discard("ldap")
        requirement_capabilities.discard("active_directory")
        requirement_capabilities.add("central_directory")

    if not requirement_documents and not requirement_capabilities:
        return None

    atom_states: dict[str, set[str]] = {}
    for document in requirement_documents:
        atom_states[f"document:{document}"] = set()
    for capability in requirement_capabilities:
        atom_states[f"feature:{capability}"] = set()

    evidence_indexes: list[int] = []
    evidence_items: list[dict[str, Any]] = []
    evidence_values: set[str] = set()
    matched_signal_names: set[str] = set()

    for final_score, company_index, _profile in scored_candidates:
        company_row = company_rows[company_index]
        company_documents = _document_ids(company_row.text)
        company_capabilities = _atomic_capability_terms(company_row.text)

        shared_documents = requirement_documents & company_documents
        shared_capabilities = requirement_capabilities & company_capabilities

        if not shared_documents and not shared_capabilities:
            continue

        support_status = _company_support_status(company_row) or "REVIEW"
        semantic_score = float(semantic_scores[company_index])
        matched_signals = sorted([*shared_documents, *shared_capabilities])

        for document in shared_documents:
            atom_states[f"document:{document}"].add(support_status)
        for capability in shared_capabilities:
            atom_states[f"feature:{capability}"].add(support_status)

        matched_signal_names.update(matched_signals)
        evidence_values.update(_critical_values(company_row.text))
        evidence_indexes.append(company_index)
        evidence_items.append(
            {
                "company_row": company_row.row_number,
                "source_filename": company_row.source_filename,
                "source_page": company_row.source_page,
                "source_kind": company_row.source_kind,
                "text": company_row.text,
                "score": round(semantic_score, 4),
                "match_score": round(float(final_score), 4),
                "support_status": support_status,
                "matched_signals": matched_signals,
                "values": company_row.values,
            }
        )

    evidence_indexes = list(dict.fromkeys(evidence_indexes))

    resolved_states: dict[str, str] = {}
    for atom, states in atom_states.items():
        if not states:
            resolved_states[atom] = "MISSING"
        elif "FC" in states and "NC" in states:
            resolved_states[atom] = "REVIEW"
        elif "FC" in states:
            resolved_states[atom] = "FC"
        elif "PC" in states:
            resolved_states[atom] = "PC"
        elif "NC" in states:
            resolved_states[atom] = "NC"
        else:
            resolved_states[atom] = "REVIEW"

    requirement_values = _critical_values(requirement_text)
    critical_mismatch = (
        not _critical_values_satisfied(
            requirement_values,
            evidence_values,
        )
    )

    states = list(resolved_states.values())
    fully_supported = [atom for atom, state in resolved_states.items() if state == "FC"]
    partially_supported = [atom for atom, state in resolved_states.items() if state == "PC"]
    not_supported = [atom for atom, state in resolved_states.items() if state == "NC"]
    missing = [atom for atom, state in resolved_states.items() if state == "MISSING"]
    review = [atom for atom, state in resolved_states.items() if state == "REVIEW"]

    if states and all(state == "FC" for state in states) and not critical_mismatch:
        status = "FC"
    elif any(state in {"FC", "PC"} for state in states):
        status = "PC"
    elif review:
        status = "REVIEW"
    else:
        status = "NC"

    evidence_items.sort(
        key=lambda item: (
            len(item["matched_signals"]),
            item["match_score"],
        ),
        reverse=True,
    )

    return {
        "status": status,
        "critical_mismatch": critical_mismatch,
        "evidence_indexes": evidence_indexes,
        "evidence_items": evidence_items[:10],
        "matched_signals": sorted(matched_signal_names),
        "fully_supported_signals": fully_supported,
        "partially_supported_signals": partially_supported,
        "missing_signals": missing,
        "not_supported_signals": not_supported,
        "review_signals": review,
    }





def _guard_semantic_only_fc(
    *,
    status: str,
    requirement_kind: str,
    requirement_text: str,
    aggregate: dict[str, Any] | None,
) -> str:
    """
    Atom veya dokuman kaniti cikmayan teknik bir
    requirement yalnizca semantic benzerlik ve
    Fully Supported etiketiyle otomatik FC olamaz.
    """

    if (
        aggregate is None
        and status == "FC"
        and requirement_kind == "CAPABILITY"
        and not _document_ids(
            requirement_text
        )
        and not _atomic_capability_terms(
            requirement_text
        )
    ):
        return "REVIEW"

    return status


def _select_display_evidence(
    *,
    status: str,
    evidence_items: list[dict[str, Any]],
    candidates: list[dict[str, Any]],
) -> dict[str, Any] | None:
    """
    Kullan?c?ya g?sterilecek denetlenebilir kan?t? se?er.

    Atom/dok?man bazl? ger?ek kan?t varsa her durumda g?sterilir.

    Yaln?zca semantic benzerlikten gelen adaylar FC/PC i?in
    geri d?n?? olarak kullan?labilir.

    REVIEW/NC durumunda do?rudan teknik kan?t yoksa alakas?z
    en-y?ksek-cosine sat?r? 'Dayanak' olarak g?sterilmez.
    """

    if evidence_items:
        return evidence_items[0]

    if (
        status in {
            "FC",
            "PC",
        }
        and candidates
    ):
        return candidates[0]

    return None



def _extract_multiple_documents(
    files: list[
        tuple[
            str,
            bytes,
        ]
    ],
) -> list[ExtractedRow]:

    rows: list[ExtractedRow] = []

    for filename, content in files:

        document_rows = (
            extract_rows(
                filename=filename,
                content=content,
            )
        )

        rows.extend(
            document_rows
        )

    return rows



# PARADOKS SPECIFICATION REQUIREMENT SPLITTER

_SPECIFICATION_FREE_TEXT_KINDS = {
    "docx",
    "pdf",
    "image",
    "pptx",
    "txt",
    "md",
    "json",
    "xml",
}


def _split_specification_requirement_row(
    row: ExtractedRow,
) -> list[ExtractedRow]:
    """
    Serbest metin sartname kaynaklarinda ayni ExtractedRow
    icine dusen birden fazla bagimsiz requirement cumlesini
    ayri requirement kayitlarina boler.

    Kaynak dosya, sayfa ve kaynak turu korunur.

    Ornek:
        Availability 99.999 olmalidir. 2FA desteklenmelidir.

    ->
        Availability 99.999 olmalidir.
        2FA desteklenmelidir.

    99.999, TLS 1.3 ve benzeri teknik degerlerdeki
    noktalar whitespace ile takip edilmedigi icin bolunmez.
    """

    source_kind = (
        row.source_kind
        or ""
    ).strip().casefold()

    if (
        source_kind
        not in _SPECIFICATION_FREE_TEXT_KINDS
    ):
        return [row]

    source_text = (
        row.text
        or ""
    ).strip()

    if not source_text:
        return []

    parts = [
        part.strip()
        for part in re.split(
            r"(?<=[.!?])\s+",
            source_text,
        )
        if part.strip()
    ]

    if len(parts) <= 1:
        return [row]

    split_rows: list[
        ExtractedRow
    ] = []

    for part in parts:
        split_rows.append(
            ExtractedRow(
                row_number=row.row_number,
                values={
                    "text": part,
                },
                text=part,
                source_filename=(
                    row.source_filename
                ),
                source_page=(
                    row.source_page
                ),
                source_kind=(
                    row.source_kind
                ),
            )
        )

    return split_rows


def _prepare_specification_rows(
    rows: list[ExtractedRow],
) -> list[ExtractedRow]:
    prepared: list[
        ExtractedRow
    ] = []

    for row in rows:

        for candidate in (
            _split_specification_requirement_row(
                row
            )
        ):

            if not (
                _is_specification_requirement(
                    candidate
                )
            ):
                continue

            prepared.append(
                candidate
            )

    return prepared


def compare_documents(
    *,
    company_files: list[
        tuple[
            str,
            bytes,
        ]
    ],
    specification_files: list[
        tuple[
            str,
            bytes,
        ]
    ],
) -> dict[str, Any]:
    import numpy as np

    company_rows = [
        row
        for row in (
            _extract_multiple_documents(
                company_files
            )
        )
        if _is_company_capability_row(
            row
        )
    ]

    specification_rows = (
        _prepare_specification_rows(
            _extract_multiple_documents(
                specification_files
            )
        )
    )

    if not company_rows:
        raise ValueError(
            "Şirket özellikleri dosyasından "
            "anlamlı yetenek satırı çıkarılamadı."
        )

    if not specification_rows:
        raise ValueError(
            "Şartname dosyasından "
            "anlamlı teknik gereksinim çıkarılamadı."
        )

    embedding_service = (
        EmbeddingService()
    )

    company_embeddings = np.asarray(
        embedding_service.embed_queries(
            [
                row.text
                for row in company_rows
            ]
        ),
        dtype=np.float32,
    )

    requirement_embeddings = np.asarray(
        embedding_service.embed_queries(
            [
                _requirement_body(row)
                for row in specification_rows
            ]
        ),
        dtype=np.float32,
    )

    results: list[
        dict[str, Any]
    ] = []

    used_company_indexes: set[int] = set()

    for requirement_index, (
        requirement,
        requirement_embedding,
    ) in enumerate(
        zip(
            specification_rows,
            requirement_embeddings,
        )
    ):
        semantic_scores = _cosine_scores(
            company_embeddings,
            requirement_embedding,
        )

        scored_candidates: list[
            tuple[
                float,
                int,
                dict[str, Any],
            ]
        ] = []

        for company_index, (
            company_row,
            semantic_score,
        ) in enumerate(
            zip(
                company_rows,
                semantic_scores,
            )
        ):
            semantic_score = float(
                semantic_score
            )
            profile = _alignment_profile(
                _requirement_body(requirement),
                company_row.text,
                semantic_score,
            )
            scored_candidates.append(
                (
                    float(
                        profile[
                            "final_score"
                        ]
                    ),
                    company_index,
                    profile,
                )
            )

        scored_candidates.sort(
            key=lambda item: (
                item[0]
            ),
            reverse=True,
        )

        if not scored_candidates:
            continue

        top_candidates = (
            scored_candidates[:10]
        )

        aggregate = (
            _aggregate_requirement_evidence(
                requirement=requirement,
                company_rows=company_rows,
                semantic_scores=(
                    semantic_scores
                ),
                scored_candidates=(
                    scored_candidates
                ),
            )
        )

        (
            best_final_score,
            best_index,
            best_profile,
        ) = top_candidates[0]

        best_company_row = (
            company_rows[
                best_index
            ]
        )
        best_semantic_score = float(
            semantic_scores[
                best_index
            ]
        )
        company_status = (
            _company_support_status(
                best_company_row
            )
        )

        requirement_text = _requirement_body(
            requirement
        )
        requirement_kind = _requirement_kind(
            requirement_text
        )

        if aggregate is not None:
            status = str(
                aggregate[
                    "status"
                ]
            )
            critical_mismatch = bool(
                aggregate[
                    "critical_mismatch"
                ]
            )
        else:
            strong = bool(
                best_profile[
                    "strong"
                ]
            )
            medium = bool(
                best_profile[
                    "medium"
                ]
            )
            critical_mismatch = bool(
                best_profile[
                    "critical_mismatch"
                ]
            )

            if strong:
                if company_status == "NC":
                    status = "NC"
                elif company_status == "PC":
                    status = "PC"
                elif company_status == "FC":
                    status = (
                        "PC"
                        if critical_mismatch
                        else "FC"
                    )
                else:
                    # Şirket satırında açık destek statüsü yoksa
                    # otomatik FC vermek yerine manuel inceleme.
                    status = "REVIEW"

            elif medium:
                if company_status == "NC":
                    status = "NC"
                elif company_status in {
                    "FC",
                    "PC",
                }:
                    status = "PC"
                else:
                    status = "REVIEW"

            else:
                if (
                    best_semantic_score
                    >= 0.80
                    and float(
                        best_profile[
                            "keyword_coverage"
                        ]
                    )
                    >= 0.22
                ):
                    status = "REVIEW"
                else:
                    status = "NC"

        # Bilgi/dokümantasyon ve ticari teyit maddelerini "üründe yok"
        # diye NC'ye düşürme. Bu maddeler teklif cevabı/manuel teyit ister.
        status = _guard_semantic_only_fc(
            status=status,
            requirement_kind=requirement_kind,
            requirement_text=requirement_text,
            aggregate=aggregate,
        )

        if requirement_kind == "INFO":
            if status == "NC":
                status = "REVIEW"
            elif (
                status == "FC"
                and _information_requires_explicit_value(
                    requirement_text
                )
                and not _evidence_has_explicit_value(
                    list(
                        aggregate["evidence_items"]
                    )
                    if aggregate
                    else []
                )
            ):
                status = "REVIEW"

        elif (
            requirement_kind == "COMMERCIAL"
            and status == "NC"
        ):
            status = "REVIEW"

        # Kullanılan şirket satırlarını extra_capabilities'ten çıkar.
        if aggregate is not None:
            used_company_indexes.update(
                aggregate[
                    "evidence_indexes"
                ]
            )
        elif status in {
            "FC",
            "PC",
        }:
            used_company_indexes.add(
                best_index
            )

        candidates: list[
            dict[str, Any]
        ] = []

        for (
            final_score,
            company_index,
            profile,
        ) in top_candidates:
            company_row = (
                company_rows[
                    company_index
                ]
            )
            semantic_score = float(
                semantic_scores[
                    company_index
                ]
            )

            candidate_signals = sorted(
                [
                    *profile[
                        "shared_documents"
                    ],
                    *profile[
                        "shared_features"
                    ],
                ]
            )

            candidates.append(
                {
                    "company_row": (
                        company_row.row_number
                    ),
                    "source_filename": (
                        company_row.source_filename
                    ),
                    "source_page": (
                        company_row.source_page
                    ),
                    "source_kind": (
                        company_row.source_kind
                    ),
                    "text": (
                        company_row.text
                    ),
                    "score": round(
                        semantic_score,
                        4,
                    ),
                    "match_score": round(
                        float(
                            final_score
                        ),
                        4,
                    ),
                    "matched_signals": (
                        candidate_signals
                    ),
                    "values": (
                        company_row.values
                    ),
                }
            )

        matched_signals = sorted(
            [
                *best_profile[
                    "shared_documents"
                ],
                *best_profile[
                    "shared_features"
                ],
            ]
        )
        evidence_items: list[
            dict[str, Any]
        ] = []

        if aggregate is not None:
            matched_signals = list(
                aggregate[
                    "matched_signals"
                ]
            )
            evidence_items = list(
                aggregate[
                    "evidence_items"
                ]
            )

        if status == "FC":
            explanation = (
                "Şartname maddesindeki tanımlanabilen teknik "
                "zorunlulukların tamamı şirket özelliklerinde "
                "desteklenen kanıtlarla karşılandı."
            )
        elif status == "PC":
            if critical_mismatch:
                explanation = (
                    "İlgili yeteneklerin bir bölümü mevcut; ancak "
                    "şartnamedeki nicel/sürüm koşullarının tamamı "
                    "birebir doğrulanmıyor."
                )
            elif aggregate is not None:
                missing_count = len(
                    aggregate[
                        "missing_signals"
                    ]
                )
                not_supported_count = len(
                    aggregate[
                        "not_supported_signals"
                    ]
                )
                explanation = (
                    "Şartname maddesi kısmen karşılanıyor. "
                    f"Doğrulanamayan atom: {missing_count}, "
                    f"desteklenmeyen atom: {not_supported_count}."
                )
            else:
                explanation = (
                    "İlgili yetenek mevcut görünüyor; ancak "
                    "şartname maddesinin tamamını doğrulayan "
                    "teknik kanıt kısmi."
                )
        elif status == "REVIEW":
            if requirement_kind == "INFO":
                explanation = (
                    "Bu madde doğrudan bir ürün özelliğinden çok "
                    "tedarikçi açıklaması, liste, doküman veya değer "
                    "talep ediyor. Mevcut özellik dosyası kesin cevap "
                    "vermiyor; teklif cevabında bilgi/teyit girilmeli."
                )
            elif requirement_kind == "COMMERCIAL":
                explanation = (
                    "Bu madde ticari/lisans/tedarik teyidi gerektiriyor. "
                    "Teknik özellik listesinden kesin uygunluk kararı "
                    "verilemez; teklif veya sözleşme tarafında doğrulanmalı."
                )
            else:
                explanation = (
                    "Benzer bir şirket özelliği bulundu; ancak "
                    "otomatik uygunluk kararı için kanıt yeterince "
                    "kesin değil. Manuel doğrulama gerekli."
                )
        else:
            if (
                aggregate is not None
                and aggregate[
                    "not_supported_signals"
                ]
            ):
                explanation = (
                    "Şirket özelliklerinde ilgili teknik yetenek "
                    "açıkça desteklenmiyor."
                )
            else:
                explanation = (
                    "Şirket özelliklerinde bu gereksinimi "
                    "doğrulayacak yeterli ve güvenilir bir "
                    "karşılık bulunamadı; edinim veya geliştirme "
                    "gerekebilir."
                )

        results.append(
            {
                "requirement_index": (
                    requirement_index + 1
                ),
                "specification_row": (
                    requirement.row_number
                ),
                "specification_source_filename": (
                    requirement.source_filename
                ),
                "specification_source_page": (
                    requirement.source_page
                ),
                "specification_source_kind": (
                    requirement.source_kind
                ),
                "requirement": (
                    requirement.text
                ),
                "requirement_values": (
                    requirement.values
                ),
                "status": status,
                "requirement_kind": (
                    requirement_kind
                ),
                "best_score": round(
                    best_semantic_score,
                    4,
                ),
                "match_score": round(
                    float(
                        best_final_score
                    ),
                    4,
                ),
                "matched_signals": (
                    matched_signals
                ),
                "evidence_items": (
                    evidence_items
                ),
                "fully_supported_signals": (
                    aggregate[
                        "fully_supported_signals"
                    ]
                    if aggregate
                    else []
                ),
                "partially_supported_signals": (
                    aggregate[
                        "partially_supported_signals"
                    ]
                    if aggregate
                    else []
                ),
                "missing_signals": (
                    aggregate[
                        "missing_signals"
                    ]
                    if aggregate
                    else []
                ),
                "not_supported_signals": (
                    aggregate[
                        "not_supported_signals"
                    ]
                    if aggregate
                    else []
                ),
                "review_signals": (
                    aggregate[
                        "review_signals"
                    ]
                    if aggregate
                    else []
                ),
                "evidence": (
                    _select_display_evidence(
                        status=status,
                        evidence_items=evidence_items,
                        candidates=candidates,
                    )
                ),
                "candidates": (
                    candidates
                ),
                "explanation": (
                    explanation
                ),
            }
        )

    extra_capabilities: list[
        dict[str, Any]
    ] = []

    for index, company_row in enumerate(
        company_rows
    ):
        if index in used_company_indexes:
            continue

        support_status = (
            _company_support_status(
                company_row
            )
        )

        # "Bizde bulunan fazla özellik" yalnız gerçekten
        # desteklenen/kısmen desteklenen satırlardan oluşsun.
        if support_status not in {
            "FC",
            "PC",
        }:
            continue

        if not _is_company_capability_row(
            company_row
        ):
            continue

        extra_capabilities.append(
            {
                "company_row": (
                    company_row.row_number
                ),
                "source_filename": (
                    company_row.source_filename
                ),
                "source_page": (
                    company_row.source_page
                ),
                "source_kind": (
                    company_row.source_kind
                ),
                "capability": (
                    company_row.text
                ),
                "values": (
                    company_row.values
                ),
            }
        )

    counts = {
        "FC": 0,
        "PC": 0,
        "NC": 0,
        "REVIEW": 0,
    }

    for result in results:
        counts[
            result[
                "status"
            ]
        ] += 1

    total = len(results)

    coverage = (
        (
            counts["FC"]
            + 0.5
            * counts["PC"]
        )
        / total
        * 100
        if total
        else 0.0
    )

    return {
        "summary": {
            "total_requirements": total,
            "fully_compliant": (
                counts["FC"]
            ),
            "partially_compliant": (
                counts["PC"]
            ),
            "non_compliant": (
                counts["NC"]
            ),
            "manual_review": (
                counts["REVIEW"]
            ),
            "coverage_percent": round(
                coverage,
                1,
            ),
        },
        "results": results,
        "extra_capabilities": (
            extra_capabilities
        ),
        "company_row_count": (
            len(company_rows)
        ),
        "specification_row_count": (
            len(specification_rows)
        ),
    }